import openpyxl
file = "openpyxl_practice.xlsx"
workbook = openpyxl.load_workbook(file)
sheet= workbook["Employees"]
rows = sheet.max_row # count numbers of rows in a excel sheet
cols = sheet.max_column # count numbers of cols in a excel sheet

#Reading all the rows & columns from excel sheet
for r in range(1,rows+1):
    for c in range(1,cols+1):
        print(sheet.cell(r,c).value,end='     ')
    print()